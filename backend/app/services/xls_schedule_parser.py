"""Детерминированный парсер расписания из Excel-таблиц (.xls/.xlsx) в формате
филиала МЭИ (Смоленск): строки-дни, колонки-группы, ячейки «лк/лб/у/кр Фамилия И.О., ауд».

Структура листа:
- строка 1: коды групп (Э - 25, ИВТ1 - 25, ...) каждые ~3 колонки
- колонка 0: дни недели (понедельник...), колонка 1: слоты времени («8.30 - 10.00»)
- ячейки занятий: «лк Физика  доц. Быков А.А.  521» — тип, предмет, звание+ФИО,
  недели («3,7,11,15 н.»), аудитория; несколько занятий в одной ячейке разделены
  большими пробелами/переносами строк; справа от группы часто «overflow»-колонка
  с параллельным занятием для другого подгруппового потока.

Выход — тот же формат items, что и у AI-парсинга parse_schedule_all, поэтому
раздача по преподавателям (teacher_matcher) работает без изменений.
"""
import io
import re
from typing import Optional

# ── Константы формата ─────────────────────────────────────────────

DAYS = {
    "понедельник": 0, "вторник": 1, "среда": 2, "четверг": 3,
    "пятница": 4, "суббота": 5, "воскресенье": 6,
}

SLOT_RE = re.compile(r"(\d{1,2})[.](\d{2})\s*[-–—]\s*(\d{1,2})[.](\d{2})")

# Код группы: «Э - 25», «ИВТ1 - 25», «ПГЭС озу - 25»
GROUP_RE = re.compile(r"^[А-ЯЁ][А-ЯЁа-яё0-9 ]{0,10}\s*[-–—]\s*2\d\s*$")

# Маркеры типа занятия в начале записи
MARKER_RE = re.compile(r"^(лк|лб|лаб|кр|кп|пз|сем|у)(?=[\s\-])", re.IGNORECASE)

MARKER_LABEL = {
    "лк": "лекция", "лаб": "лабораторная", "лб": "лабораторная",
    "у": "практика", "кр": "курсовая работа", "кп": "курсовой проект",
    "пз": "проектное занятие", "сем": "семинар",
}

# Звания/должности перед ФИО
RANK = (r"(?:доц|проф|ассистент|асс|ст\.?\s?преп|ст\.?\s?пр|вн|зав|акад"
        r"|к\.т\.н|д\.т\.н|доцент|профессор|старший\s+преподаватель)")

# ФИО: ВСЕГДА со званием — «доц. Быков А.А.», «асс.Жарков А.П.», «ст.пр. Гаврилов А.И».
# Звание обязательно: иначе подписи («составила Космачева О.») и прочие
# фамилии без контекста занятия превращаются в преподавателей.
TEACHER_RE = re.compile(
    rf"(?<![а-яё])(?:{RANK})[.,]?\s*([А-ЯЁ][а-яё]{{2,}})\s+"
    rf"(?:([А-ЯЁ])[.,]\s*([А-ЯЁ]?)[.,]?|([А-ЯЁ][а-яё]{{2,}})\s+([А-ЯЁ])[.,])"
)

WEEK_LIST_RE = re.compile(r"\b\d{1,2}(?:\s*,\s*\d{1,2})+\s*н\.?\b")
WEEK_SPAN_RE = re.compile(r"\b(?:со|с)\s+\d+\s+по\s+\d+\s+нед\.?\b", re.IGNORECASE)
PARA_RE = re.compile(r"\b\d+\s*и\s*\d+\s*(?:пара|пар)\b")
PGR_RE = re.compile(r"\b\d+\s*пгр\.?\b")

# Аудитории: «А 214», «В 01», «521 А», «425»
ROOM_LET_RE = re.compile(r"\b([А-В])\s+(\d{1,4})\b")
ROOM_NUM_RE = re.compile(r"\b(\d{2,4})\s*([А-В])?\b")

# Занятия без преподавателя — всегда отдельной записью (и выбрасываются при финализации)
NO_TEACHER_SUBJECTS = ("элективные курсы",)


def _cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float):
        if value == int(value):
            return str(int(value))
        return str(value)
    return str(value).replace("\r", "").strip()


def _norm_group(txt: str) -> str:
    s = re.sub(r"\s*[-–—]\s*", "-", txt)
    return re.sub(r"\s+", "", s)


# ── Чтение таблиц ─────────────────────────────────────────────────

def _load_grids(file_bytes: bytes, filename: str) -> list[tuple[str, list[list[str]]]]:
    ext = ("." + filename.rsplit(".", 1)[-1]).lower() if "." in filename else ""
    grids = []
    if ext == ".xls":
        import xlrd
        book = xlrd.open_workbook(file_contents=file_bytes)
        for sh in book.sheets():
            grid = [[_cell_text(sh.cell_value(r, c)) for c in range(sh.ncols)]
                    for r in range(sh.nrows)]
            grids.append((sh.name, grid))
    elif ext == ".xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(file_bytes), data_only=True)
        for ws in wb.worksheets:
            grid = [[_cell_text(v) for v in row] for row in ws.iter_rows(values_only=True)]
            grids.append((ws.title, grid))
    return grids


# ── Разбор одной ячейки-блока на записи ───────────────────────────

class _Record:
    __slots__ = ("parts",)

    def __init__(self, part: str = ""):
        self.parts = [part] if part else []


def _split_records(cell_texts: list[str]) -> list[_Record]:
    """Сегменты ячеек склеиваются в записи: новый маркер (лк/лб/у/кр/...) начинает запись.

    Также новую запись начинает сегмент вида «Химия доц. Короткова Г.В.» —
    SUBJECT с заглавной буквы перед ФИО, когда текущая запись уже «закрыта»
    преподавателем. Переносы слов («строения проф. Якименко И.В.») склеиваются.
    """
    records: list[_Record] = []
    cur: Optional[_Record] = None
    for text in cell_texts:
        for line in text.split("\n"):
            for seg in re.split(r" {3,}", line):
                seg = seg.strip()
                if not seg:
                    continue
                low = seg.lower()
                starts_new = bool(MARKER_RE.match(seg)) or low.startswith(NO_TEACHER_SUBJECTS)
                if not starts_new and cur is not None and cur.parts:
                    cur_raw = " ".join(cur.parts)
                    if TEACHER_RE.search(cur_raw):
                        # у текущей записи уже есть преподаватель — возможен новый предмет
                        m = TEACHER_RE.search(seg)
                        pre = seg[: m.start()].strip() if m else ""
                        if pre and pre[0].isupper() and 2 <= len(pre) <= 40:
                            starts_new = True
                if starts_new or cur is None:
                    cur = _Record(seg)
                    records.append(cur)
                else:
                    cur.parts.append(seg)
    return [r for r in records if r.parts]


def _extract_teachers(raw: str) -> list[tuple[str, int, int]]:
    """[(ФИО, начало, конец)] всех преподавателей в записи."""
    out = []
    for m in TEACHER_RE.finditer(raw):
        sur = m.group(1)
        if m.group(2):
            inits = m.group(2) + "." + (m.group(3) + "." if m.group(3) else "")
        else:
            inits = m.group(4) + " " + (m.group(5) + "." if m.group(5) else "")
        name = f"{sur} {inits}".replace("  ", " ").strip()
        # Защита от мусора: «В.В. Рожков» (инициалы перед фамилией в шапках) не ловим —
        # там нет паттерна «фамилия + инициал».
        out.append((re.sub(r"\s+", " ", name), m.start(), m.end()))
    return out


def _finalize_record(raw: str) -> Optional[dict]:
    """Запись текста → {subject, marker, teachers, weeks, rooms} или None."""
    teachers = _extract_teachers(raw)
    if not teachers:
        return None

    # Недели: «3,7,11,15 н.» / «со 2 по 12 нед.» → «3,7,11,15» / «2 по 12»
    weeks = []
    for rx in (WEEK_LIST_RE, WEEK_SPAN_RE):
        for m in rx.finditer(raw):
            w = re.sub(r"\s+", " ", m.group(0)).strip(" .,")
            w = re.sub(r"\s*(?:нед|н)\.?\s*$", "", w, flags=re.IGNORECASE)
            w = re.sub(r"^(?:со|с)\s+", "", w, flags=re.IGNORECASE)
            if w and w not in weeks:
                weeks.append(w)

    # Текст без преподавателей/недель/пара/пгр — источник аудиторий
    cleaned = raw
    spans = [t[1:3] for t in teachers]
    for rx in (WEEK_LIST_RE, WEEK_SPAN_RE, PARA_RE, PGR_RE):
        for m in rx.finditer(cleaned):
            spans.append((m.start(), m.end()))
    spans = sorted(set(spans), reverse=True)
    for a, b in spans:
        a, b = min(a, len(cleaned)), min(b, len(cleaned))
        cleaned = cleaned[:a] + " " + cleaned[b:]

    rooms: list[str] = []
    no_rooms = cleaned
    for m in ROOM_LET_RE.finditer(no_rooms):
        r = f"{m.group(1)} {m.group(2)}"
        if r not in rooms:
            rooms.append(r)
    tmp = ROOM_LET_RE.sub(" ", no_rooms)
    for m in ROOM_NUM_RE.finditer(tmp):
        r = m.group(1) + (f" {m.group(2)}" if m.group(2) else "")
        if r not in rooms:
            rooms.append(r)

    # Предмет: до первого преподавателя
    subject = raw[: teachers[0][1]].strip(" ,.;:-")
    marker = None
    mm = MARKER_RE.match(subject)
    if mm:
        marker = MARKER_LABEL.get(mm.group(1).lower())
        subject = subject[mm.end():]
    # Собираем переносы вида «машинострои-тельным» в нормальные слова
    subject = re.sub(r"(?<=[а-яё])--?(?=[а-яё])", "", subject)
    subject = re.sub(r"\s{2,}", " ", subject).strip(" ,.;:-")
    subject = re.sub(r"\s+", " ", subject)
    if not subject or len(subject) < 3:
        return None

    return {
        "subject": subject,
        "marker": marker,
        "teachers": [t[0] for t in teachers],
        "weeks": weeks,
        "rooms": rooms[:4],
    }


# ── Разбор листа ──────────────────────────────────────────────────

def _parse_sheet(name: str, grid: list[list[str]]) -> list[dict]:
    if len(grid) < 5:
        return []

    ncols = max((len(r) for r in grid), default=0)
    grid = [row + [""] * (ncols - len(row)) for row in grid]

    # Строка с группами — одна из первых трёх
    groups: list[tuple[int, str]] = []
    header_row = -1
    for ri in range(min(3, len(grid))):
        found = [(ci, _norm_group(_cell_text(cv))) for ci, cv in enumerate(grid[ri])
                 if _cell_text(cv) and GROUP_RE.match(_cell_text(cv))]
        if len(found) > len(groups):
            groups, header_row = found, ri
    if len(groups) < 2:
        return []

    # Диапазон колонок каждой группы: от её колонки до следующей группы
    ranges = []
    for i, (gc, gname) in enumerate(groups):
        end = groups[i + 1][0] - 1 if i + 1 < len(groups) else ncols - 1
        end = min(end, gc + 3)
        ranges.append((gc, end, gname))

    # Слоты: строки с днём (col 0) и временем (col 1)
    slots: list[dict] = []
    cur_day: Optional[int] = None
    for ri in range(header_row + 1, len(grid)):
        row = grid[ri]
        d = _cell_text(row[0]).lower().strip()
        if d in DAYS:
            cur_day = DAYS[d]
        m = SLOT_RE.search(_cell_text(row[1]))
        if m and cur_day is not None:
            slots.append({
                "row": ri, "day": cur_day,
                "start": f"{int(m.group(1)):02d}:{m.group(2)}",
                "end": f"{int(m.group(3)):02d}:{m.group(4)}",
            })
    if not slots:
        return []
    for i, s in enumerate(slots):
        s["row_end"] = (slots[i + 1]["row"] if i + 1 < len(slots) else len(grid)) - 1

    items: list[dict] = []
    for s in slots:
        for lo, hi, gname in ranges:
            texts = []
            for ri in range(s["row"], s["row_end"] + 1):
                for ci in range(lo, hi + 1):
                    t = _cell_text(grid[ri][ci])
                    if t:
                        texts.append(t)
            if not texts:
                continue
            for rec in _split_records(texts):
                fin = _finalize_record(" ".join(rec.parts))
                if not fin:
                    continue
                title = fin["subject"]
                extras = []
                if fin["marker"]:
                    extras.append(fin["marker"])
                if fin["weeks"]:
                    extras.append("нед. " + "; ".join(
                        re.sub(r"\s*,\s*", ",", w) for w in fin["weeks"]))
                if extras:
                    title = f"{title} ({', '.join(extras)})"
                for teacher in fin["teachers"]:
                    items.append({
                        "title": title,
                        "day_of_week": s["day"],
                        "start_time": s["start"],
                        "end_time": s["end"],
                        "group_name": gname,
                        "room": " / ".join(fin["rooms"]),
                        "type": "lesson",
                        "teacher": teacher,
                        "source": "xls_import",
                        "sheet": name,
                    })
    return items


# ── Публичное API ─────────────────────────────────────────────────

def parse_xls_schedule(file_bytes: bytes, filename: str) -> list[dict]:
    """Парсит Excel-расписание в items-формат AI-парсинга.

    Возвращает [] если файл не похож на расписание — вызывающий может
    откатиться на текстовый извлечение + AI.
    """
    items: list[dict] = []
    try:
        grids = _load_grids(file_bytes, filename)
    except Exception:
        return []
    for name, grid in grids:
        try:
            items.extend(_parse_sheet(name, grid))
        except Exception:
            continue

    # Дедупликация внутри файла (один преподаватель, одна пара продублирована)
    seen = set()
    uniq = []
    for it in items:
        key = (it["teacher"].lower(), it["day_of_week"], it["start_time"],
               it["title"].lower(), it["group_name"])
        if key in seen:
            continue
        seen.add(key)
        uniq.append(it)
    return uniq


def extract_xls_text(file_bytes: bytes, filename: str) -> str:
    """Запасной текстовый дамп таблиц (для AI-парсинга, если структурный не сработал)."""
    lines: list[str] = []
    for name, grid in _load_grids(file_bytes, filename):
        lines.append(f"=== Лист: {name} ===")
        for row in grid:
            cells = [c.replace("\n", " ").strip() for c in row]
            cells = [c for c in cells if c]
            if cells:
                lines.append(" | ".join(cells))
    return "\n".join(lines)
