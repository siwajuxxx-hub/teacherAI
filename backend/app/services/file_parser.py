"""Парсинг файлов: извлечение текста из PDF, DOCX, TXT и изображений."""
import io
import os
from typing import Optional


async def extract_text_from_file(file_bytes: bytes, filename: str) -> str:
    """Извлекает текст из файла в зависимости от расширения."""
    ext = os.path.splitext(filename)[1].lower()

    if ext == ".txt":
        return file_bytes.decode("utf-8", errors="ignore")

    elif ext == ".pdf":
        return _extract_pdf(file_bytes)

    elif ext in (".docx", ".doc"):
        return _extract_docx(file_bytes)

    elif ext in (".png", ".jpg", ".jpeg"):
        # Для изображений текст не извлекается — нужно отправлять в AI с vision
        return "[IMAGE] Изображение требует обработки через AI с поддержкой vision"

    else:
        raise ValueError(f"Неподдерживаемый формат файла: {ext}")


def _extract_pdf(file_bytes: bytes) -> str:
    try:
        import pdfplumber
        text_parts = []
        with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
            for page in pdf.pages:
                text = page.extract_text()
                if text:
                    text_parts.append(text)

        if not text_parts:
            # Пробуем извлечь таблицы
            with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
                for page in pdf.pages:
                    tables = page.extract_tables()
                    for table in tables:
                        for row in table:
                            text_parts.append(" | ".join([str(cell) if cell else "" for cell in row]))

        return "\n\n".join(text_parts) if text_parts else ""
    except Exception as e:
        raise Exception(f"Ошибка парсинга PDF: {str(e)}")


def _extract_docx(file_bytes: bytes) -> str:
    try:
        from docx import Document
        doc = Document(io.BytesIO(file_bytes))
        text_parts = []

        # Извлекаем текст из параграфов
        for para in doc.paragraphs:
            if para.text.strip():
                text_parts.append(para.text.strip())

        # Извлекаем таблицы
        for table in doc.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                text_parts.append(" | ".join(cells))

        return "\n".join(text_parts) if text_parts else ""
    except Exception as e:
        raise Exception(f"Ошибка парсинга DOCX: {str(e)}")


async def extract_text_from_image(file_bytes: bytes) -> Optional[str]:
    """Заглушка — реальное извлечение текста из изображения делается через AI с vision."""
    return None  # Возвращает None, чтобы вызывающий код отправил изображение в AI